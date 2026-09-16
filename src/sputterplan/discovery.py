"""Profile discovery and first-run configuration helpers."""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .errors import ConfigurationError, OutputError, ProfileError


@dataclass(frozen=True)
class ProfileInspection:
    path: Path
    format: str
    sheet_names: tuple[str, ...]
    selected_sheet: str | None
    header_row: int
    headers: tuple[str, ...]
    preview_rows: tuple[tuple[Any, ...], ...]
    warnings: tuple[str, ...] = ()


def _display_header(value: Any, index: int) -> str:
    text = "" if value is None else str(value).strip()
    return text or f"column_{index + 1}"


def _header_warnings(headers: tuple[str, ...]) -> tuple[str, ...]:
    duplicates = sorted({name for name in headers if headers.count(name) > 1})
    if not duplicates:
        return ()
    return (
        "Duplicate column names detected: "
        + ", ".join(duplicates)
        + ". Select those columns by number (or Excel letter for workbooks) in the configuration.",
    )


def inspect_profile_file(
    path: str | Path,
    *,
    sheet: str | None = None,
    header_row: int = 1,
    preview_count: int = 5,
) -> ProfileInspection:
    """Read headers and a small preview without changing the source file."""
    profile_path = Path(path).expanduser().resolve()
    if not profile_path.exists():
        raise ProfileError(f"Profile file does not exist: {profile_path}")
    if header_row < 1:
        raise ProfileError("header_row is one-based and must be at least 1.")
    if preview_count < 0:
        raise ProfileError("preview_count cannot be negative.")

    suffix = profile_path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            with profile_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, delimiter=delimiter))
        except OSError as exc:
            raise ProfileError(f"Could not read {profile_path}: {exc}") from exc
        if header_row > len(rows):
            raise ProfileError(f"header_row {header_row} is beyond the end of {profile_path.name}.")
        raw_headers = rows[header_row - 1]
        preview = rows[header_row : header_row + preview_count]
        headers = tuple(_display_header(value, index) for index, value in enumerate(raw_headers))
        return ProfileInspection(
            path=profile_path,
            format=suffix.lstrip("."),
            sheet_names=(),
            selected_sheet=None,
            header_row=header_row,
            headers=headers,
            preview_rows=tuple(tuple(row) for row in preview),
            warnings=_header_warnings(headers),
        )

    if suffix not in {".xlsx", ".xlsm"}:
        raise ProfileError(f"Unsupported profile format {suffix!r}; use CSV, TSV, XLSX, or XLSM.")
    try:
        workbook = load_workbook(profile_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ProfileError(f"Could not open Excel profile {profile_path}: {exc}") from exc
    try:
        sheet_names = tuple(workbook.sheetnames)
        if sheet is not None and sheet not in sheet_names:
            raise ProfileError(
                f"Worksheet {sheet!r} not found. Available sheets: {', '.join(sheet_names)}"
            )
        selected = sheet or sheet_names[0]
        warnings: list[str] = []
        if sheet is None and len(sheet_names) > 1:
            warnings.append(
                f"Workbook has {len(sheet_names)} sheets; previewing the first ({selected!r}). "
                "Use --sheet to select another."
            )
        worksheet = workbook[selected]
        if header_row > worksheet.max_row:
            raise ProfileError(
                f"header_row {header_row} is beyond the end of worksheet {selected!r}."
            )
        rows = list(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=min(worksheet.max_row, header_row + preview_count),
                values_only=True,
            )
        )
        raw_headers = rows[0]
        headers = tuple(_display_header(value, index) for index, value in enumerate(raw_headers))
        warnings.extend(_header_warnings(headers))
        return ProfileInspection(
            path=profile_path,
            format=suffix.lstrip("."),
            sheet_names=sheet_names,
            selected_sheet=selected,
            header_row=header_row,
            headers=headers,
            preview_rows=tuple(tuple(row) for row in rows[1:]),
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def parse_target_spec(spec: str) -> tuple[str, float]:
    """Parse CLI target syntax NAME=RATE."""
    if "=" not in spec:
        raise ConfigurationError(
            f"Invalid --target {spec!r}; use NAME=RATE, for example Ni=0.1003."
        )
    name, raw_rate = (part.strip() for part in spec.split("=", 1))
    if not name:
        raise ConfigurationError("Target name cannot be empty.")
    try:
        rate = float(raw_rate)
    except ValueError as exc:
        raise ConfigurationError(
            f"Target {name!r} rate must be a number, got {raw_rate!r}."
        ) from exc
    if rate <= 0:
        raise ConfigurationError(f"Target {name!r} rate must be positive.")
    return name, rate


def parse_column_mapping(spec: str) -> tuple[str, str | int]:
    """Parse CLI column mapping syntax NAME=COLUMN."""
    if "=" not in spec:
        raise ConfigurationError(
            f"Invalid column mapping {spec!r}; use NAME=COLUMN, for example Ni=B."
        )
    name, raw_column = (part.strip() for part in spec.split("=", 1))
    if not name or not raw_column:
        raise ConfigurationError(f"Invalid column mapping {spec!r}; both values are required.")
    column: str | int = int(raw_column) if raw_column.isdigit() else raw_column
    if isinstance(column, int) and column < 1:
        raise ConfigurationError("Column numbers are one-based and must be at least 1.")
    return name, column


def _validate_column_spec(spec: str | int, inspection: ProfileInspection, label: str) -> str | int:
    if isinstance(spec, int):
        if 1 <= spec <= len(inspection.headers):
            return spec
    elif spec in inspection.headers:
        if inspection.headers.count(spec) > 1:
            raise ConfigurationError(
                f"{label} column name {spec!r} is duplicated. Select it by one-based number"
                + (" or Excel letter." if inspection.format in {"xlsx", "xlsm"} else ".")
            )
        return spec
    elif inspection.format in {"xlsx", "xlsm"} and spec.isalpha() and spec.isupper():
        try:
            if column_index_from_string(spec) <= len(inspection.headers):
                return spec
        except ValueError:
            pass
    raise ConfigurationError(
        f"{label} column {spec!r} was not found. Available columns: "
        + ", ".join(inspection.headers)
    )


def write_starter_config(
    profile_path: str | Path,
    output_path: str | Path,
    target_specs: list[str],
    *,
    sheet: str | None = None,
    header_row: int = 1,
    distance_column: str | int | None = None,
    composition_basis: str = "fraction",
    composition_sum_tolerance: float = 1e-6,
    normalize_compositions: bool = False,
    total_power_w: float = 300.0,
    off_below_power_w: float = 0.0,
    min_stable_power_w: float = 0.0,
    composition_columns: dict[str, str | int] | None = None,
    overwrite: bool = False,
) -> Path:
    """Create a valid starter YAML configuration from detected profile headers."""
    inspection = inspect_profile_file(profile_path, sheet=sheet, header_row=header_row)
    if not target_specs:
        raise ConfigurationError(
            "At least one --target NAME=RATE is required. Run inspect-profile to see column names."
        )
    targets = [parse_target_spec(spec) for spec in target_specs]
    folded = [name.casefold() for name, _ in targets]
    if len(set(folded)) != len(folded):
        raise ConfigurationError("Each --target name must be unique, ignoring capitalization.")
    mappings = composition_columns or {}
    unknown_mappings = set(mappings) - {name for name, _ in targets}
    if unknown_mappings:
        raise ConfigurationError(
            "Column mappings name unconfigured targets: " + ", ".join(sorted(unknown_mappings))
        )
    missing = [
        name for name, _ in targets if name not in inspection.headers and name not in mappings
    ]
    if missing:
        raise ConfigurationError(
            "Target columns not found in the selected header row: " + ", ".join(missing)
        )
    selected_distance: str | int = distance_column or inspection.headers[0]
    selected_distance = _validate_column_spec(selected_distance, inspection, "Distance")
    selected_compositions = {
        name: _validate_column_spec(mappings.get(name, name), inspection, f"Composition for {name}")
        for name, _ in targets
    }
    if composition_basis not in {"fraction", "percent"}:
        raise ConfigurationError("composition_basis must be 'fraction' or 'percent'.")
    if composition_sum_tolerance <= 0:
        raise ConfigurationError("composition_sum_tolerance must be positive.")
    if total_power_w <= 0:
        raise ConfigurationError("total_power_w must be positive.")
    if off_below_power_w < 0 or min_stable_power_w < off_below_power_w:
        raise ConfigurationError(
            "Power thresholds require 0 <= off_below_power_w <= min_stable_power_w."
        )

    destination = Path(output_path).expanduser().resolve()
    if destination.exists() and not overwrite:
        raise OutputError(
            f"Configuration already exists: {destination}. Choose another path or use --force."
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    relative_profile = Path(os.path.relpath(inspection.path, destination.parent)).as_posix()
    profile_data: dict[str, Any] = {
        "path": relative_profile,
        "header_row": header_row,
        "distance_column": selected_distance,
        "composition_basis": composition_basis,
        "composition_columns": selected_compositions,
        "composition_sum_tolerance": composition_sum_tolerance,
        "normalize_compositions": normalize_compositions,
    }
    if inspection.selected_sheet is not None:
        profile_data["sheet"] = inspection.selected_sheet
    data = {
        "name": f"{inspection.path.stem} sputtering plan",
        "profile": profile_data,
        "targets": [
            {
                "name": name,
                "rate_nm_per_min_per_watt": rate,
                "off_below_power_w": off_below_power_w,
                "min_stable_power_w": min_stable_power_w,
            }
            for name, rate in targets
        ],
        "operation": {
            "strategy": "total_power",
            "total_power_w": total_power_w,
            "hardware_limit_policy": "clip",
        },
        "planner": {
            "power_tolerance_abs_w": 2.0,
            "power_tolerance_rel": 0.05,
            "composition_tolerance_abs": 0.02,
            "max_breakpoints": 100,
            "ramp_rate_resolution_w_per_s": 0.001,
        },
    }
    destination.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    return destination
