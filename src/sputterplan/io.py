"""Composition-profile readers for CSV, TSV, and Excel workbooks."""

from __future__ import annotations

import csv
from collections.abc import Sequence
from typing import Any

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .configuration import PlanConfig, ProfileConfig
from .errors import ProfileError
from .results import CompositionProfile


def _as_float(value: Any, row: int, column: str) -> float:
    if value is None or value == "":
        raise ProfileError(f"Missing value at row {row}, column {column!r}.")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ProfileError(
            f"Expected a number at row {row}, column {column!r}; got {value!r}."
        ) from exc
    if not np.isfinite(result):
        raise ProfileError(f"Non-finite value at row {row}, column {column!r}.")
    return result


def _resolve_header_index(headers: Sequence[Any], spec: str | int, label: str) -> int:
    if isinstance(spec, int):
        if spec < 1 or spec > len(headers):
            raise ProfileError(f"{label} column number {spec} is outside the input table.")
        return spec - 1
    matches = [index for index, value in enumerate(headers) if str(value).strip() == spec]
    if not matches:
        raise ProfileError(f"Could not find {label} column {spec!r}.")
    if len(matches) > 1:
        raise ProfileError(
            f"Column label {spec!r} is duplicated; select it by one-based column number."
        )
    return matches[0]


def _excel_column_index(headers: Sequence[Any], spec: str | int, label: str) -> int:
    if isinstance(spec, str) and spec.isalpha() and len(spec) <= 3:
        index = column_index_from_string(spec.upper()) - 1
        if index >= len(headers):
            raise ProfileError(f"{label} column {spec!r} is outside the worksheet.")
        return index
    return _resolve_header_index(headers, spec, label)


def _column_specs(profile: ProfileConfig, target_names: Sequence[str]) -> list[str | int]:
    return [profile.composition_columns.get(name, name) for name in target_names]


def _read_delimited(
    profile: ProfileConfig, target_names: Sequence[str], delimiter: str
) -> tuple[np.ndarray, np.ndarray]:
    try:
        handle = profile.path.open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise ProfileError(f"Could not open {profile.path}: {exc}") from exc
    with handle:
        rows = list(csv.reader(handle, delimiter=delimiter))
    if not rows:
        raise ProfileError(f"Composition profile is empty: {profile.path}")
    header_index = profile.header_row - 1
    if header_index >= len(rows):
        raise ProfileError("profile.header_row is beyond the end of the file.")
    headers = rows[header_index]
    distance_index = _resolve_header_index(headers, profile.distance_column, "distance")
    composition_indices = [
        _resolve_header_index(headers, spec, f"composition for {name}")
        for name, spec in zip(target_names, _column_specs(profile, target_names))
    ]
    start = (profile.data_start_row - 1) if profile.data_start_row else header_index + 1
    stop = profile.data_end_row if profile.data_end_row else len(rows)
    distances: list[float] = []
    compositions: list[list[float]] = []
    for row_number, row in enumerate(rows[start:stop], start=start + 1):
        if not row or all(value.strip() == "" for value in row):
            continue
        padded = row + [""] * max(0, len(headers) - len(row))
        distances.append(
            _as_float(padded[distance_index], row_number, str(profile.distance_column))
        )
        compositions.append(
            [
                _as_float(padded[index], row_number, str(spec))
                for index, spec in zip(composition_indices, _column_specs(profile, target_names))
            ]
        )
    return np.asarray(distances, dtype=float), np.asarray(compositions, dtype=float)


def _read_excel(
    profile: ProfileConfig, target_names: Sequence[str]
) -> tuple[np.ndarray, np.ndarray]:
    try:
        workbook = load_workbook(profile.path, read_only=True, data_only=True)
    except Exception as exc:
        raise ProfileError(f"Could not open Excel profile {profile.path}: {exc}") from exc
    try:
        if profile.sheet is None:
            if len(workbook.sheetnames) != 1:
                raise ProfileError(
                    "profile.sheet is required when an Excel workbook has multiple sheets."
                )
            sheet = workbook[workbook.sheetnames[0]]
        elif profile.sheet not in workbook.sheetnames:
            raise ProfileError(
                f"Worksheet {profile.sheet!r} not found. Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
        else:
            sheet = workbook[profile.sheet]

        headers = list(
            next(
                sheet.iter_rows(
                    min_row=profile.header_row,
                    max_row=profile.header_row,
                    values_only=True,
                )
            )
        )
        distance_index = _excel_column_index(headers, profile.distance_column, "distance")
        specs = _column_specs(profile, target_names)
        composition_indices = [
            _excel_column_index(headers, spec, f"composition for {name}")
            for name, spec in zip(target_names, specs)
        ]
        start = profile.data_start_row or profile.header_row + 1
        stop = profile.data_end_row or sheet.max_row
        distances: list[float] = []
        compositions: list[list[float]] = []
        for row_number, values_tuple in enumerate(
            sheet.iter_rows(min_row=start, max_row=stop, values_only=True), start=start
        ):
            values = list(values_tuple)
            distance_value = values[distance_index] if distance_index < len(values) else None
            if distance_value is None or distance_value == "":
                continue
            distances.append(_as_float(distance_value, row_number, str(profile.distance_column)))
            compositions.append(
                [
                    _as_float(
                        values[index] if index < len(values) else None,
                        row_number,
                        str(spec),
                    )
                    for index, spec in zip(composition_indices, specs)
                ]
            )
        return np.asarray(distances, dtype=float), np.asarray(compositions, dtype=float)
    finally:
        workbook.close()


def _validate_profile(
    distance: np.ndarray,
    compositions: np.ndarray,
    target_names: Sequence[str],
    profile: ProfileConfig,
) -> CompositionProfile:
    if distance.ndim != 1 or distance.size < 2:
        raise ProfileError("The composition profile must contain at least two distance points.")
    if compositions.shape != (distance.size, len(target_names)):
        raise ProfileError("Composition table dimensions do not match the configured targets.")
    if not np.all(np.isfinite(distance)) or not np.all(np.isfinite(compositions)):
        raise ProfileError("Distance and composition values must be finite.")
    if np.any(np.diff(distance) <= 0):
        bad = int(np.flatnonzero(np.diff(distance) <= 0)[0])
        raise ProfileError(
            f"Distance must increase strictly; rows {bad + 1} and {bad + 2} are not increasing."
        )
    if np.any(compositions < 0):
        row, column = np.argwhere(compositions < 0)[0]
        raise ProfileError(
            f"Negative composition at profile row {row + 1}, target {target_names[column]!r}."
        )
    if profile.composition_basis == "percent":
        compositions = compositions / 100.0
    totals = np.sum(compositions, axis=1)
    if np.any(totals <= 0):
        row = int(np.flatnonzero(totals <= 0)[0])
        raise ProfileError(f"Composition sum is zero at profile row {row + 1}.")
    if profile.normalize_compositions:
        compositions = compositions / totals[:, None]
    else:
        error = np.abs(totals - 1.0)
        if np.any(error > profile.composition_sum_tolerance):
            row = int(np.argmax(error))
            raise ProfileError(
                f"Compositions sum to {totals[row]:.8g} at profile row {row + 1}, not 1. "
                "Fix the data or set profile.normalize_compositions: true explicitly."
            )
    return CompositionProfile(
        distance_nm=np.asarray(distance, dtype=float),
        compositions=np.asarray(compositions, dtype=float),
        target_names=tuple(target_names),
    )


def load_profile(config: PlanConfig) -> CompositionProfile:
    profile = config.profile
    suffix = profile.path.suffix.lower()
    if suffix == ".csv":
        distance, compositions = _read_delimited(profile, config.target_names, ",")
    elif suffix == ".tsv":
        distance, compositions = _read_delimited(profile, config.target_names, "\t")
    elif suffix in {".xlsx", ".xlsm"}:
        distance, compositions = _read_excel(profile, config.target_names)
    else:
        raise ProfileError(f"Unsupported profile format {suffix!r}; use CSV, TSV, XLSX, or XLSM.")
    return _validate_profile(distance, compositions, config.target_names, profile)
