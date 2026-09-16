from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from sputterplan.configuration import load_config
from sputterplan.discovery import (
    inspect_profile_file,
    parse_column_mapping,
    parse_target_spec,
    write_starter_config,
)
from sputterplan.errors import ConfigurationError, ProfileError


def test_inspect_csv_and_tsv(tmp_path: Path):
    csv_path = tmp_path / "profile.csv"
    csv_path.write_text("distance_nm,A,B\n0,0.2,0.8\n1,0.4,0.6\n", encoding="utf-8")
    inspected = inspect_profile_file(csv_path, preview_count=1)
    assert inspected.format == "csv"
    assert inspected.headers == ("distance_nm", "A", "B")
    assert inspected.preview_rows == (("0", "0.2", "0.8"),)

    tsv_path = tmp_path / "profile.tsv"
    tsv_path.write_text("distance_nm\tA\n0\t1\n1\t1\n", encoding="utf-8")
    assert inspect_profile_file(tsv_path).headers == ("distance_nm", "A")


def test_inspection_warns_about_duplicate_headers(tmp_path: Path):
    profile = tmp_path / "duplicate.csv"
    profile.write_text("distance_nm,Ni,Ni\n0,0.5,0.5\n", encoding="utf-8")
    assert "Duplicate column names" in inspect_profile_file(profile).warnings[0]


def test_inspect_excel_sheets_and_preview(tmp_path: Path):
    workbook_path = tmp_path / "profile.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Gradient"
    first.append(["ignored", "ignored"])
    first.append(["distance_nm", "Ni"])
    first.append([0, 1])
    workbook.create_sheet("Notes").append(["text"])
    workbook.save(workbook_path)

    inspected = inspect_profile_file(workbook_path, header_row=2, preview_count=1)
    assert inspected.sheet_names == ("Gradient", "Notes")
    assert inspected.selected_sheet == "Gradient"
    assert inspected.headers == ("distance_nm", "Ni")
    assert inspected.preview_rows == ((0, 1),)
    assert "previewing the first" in inspected.warnings[0]

    with pytest.raises(ProfileError, match="not found"):
        inspect_profile_file(workbook_path, sheet="Missing")


@pytest.mark.parametrize("spec, expected", [("Ni=0.1003", ("Ni", 0.1003)), (" A = 2 ", ("A", 2.0))])
def test_parse_target_spec(spec, expected):
    assert parse_target_spec(spec) == expected


@pytest.mark.parametrize("spec", ["Ni", "=0.1", "Ni=abc", "Ni=0", "Ni=-1"])
def test_parse_target_spec_errors(spec):
    with pytest.raises(ConfigurationError):
        parse_target_spec(spec)


@pytest.mark.parametrize(
    "spec, expected",
    [("Ni=B", ("Ni", "B")), ("Ti=3", ("Ti", 3)), (" Nb = Nb at.% ", ("Nb", "Nb at.%"))],
)
def test_parse_column_mapping(spec, expected):
    assert parse_column_mapping(spec) == expected


@pytest.mark.parametrize("spec", ["Ni", "=B", "Ni=", "Ni=0"])
def test_parse_column_mapping_errors(spec):
    with pytest.raises(ConfigurationError):
        parse_column_mapping(spec)


def test_write_starter_config_round_trips(tmp_path: Path):
    profile = tmp_path / "gradient.csv"
    profile.write_text("distance_nm,Ni,Ti\n0,25,75\n1,50,50\n", encoding="utf-8")
    destination = tmp_path / "config" / "plan.yaml"
    result = write_starter_config(
        profile,
        destination,
        ["Ni=0.1", "Ti=0.05"],
        composition_basis="percent",
        total_power_w=250,
        off_below_power_w=1,
        min_stable_power_w=2,
    )
    assert result == destination.resolve()
    raw = yaml.safe_load(destination.read_text(encoding="utf-8"))
    assert raw["profile"]["path"] == "../gradient.csv"
    loaded = load_config(destination)
    assert loaded.target_names == ("Ni", "Ti")
    assert loaded.operation.total_power_w == 250


def test_starter_config_supports_explicit_excel_columns(tmp_path: Path):
    workbook_path = tmp_path / "duplicate_headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gradient"
    sheet.append(["distance_nm", "Ni", "Ti", "Ni", "Ti"])
    sheet.append([0, 0.25, 0.75, 1, 0])
    sheet.append([1, 0.5, 0.5, 0, 1])
    workbook.save(workbook_path)
    destination = tmp_path / "plan.yaml"
    write_starter_config(
        workbook_path,
        destination,
        ["Ni=0.1", "Ti=0.2"],
        sheet="Gradient",
        distance_column="A",
        composition_columns={"Ni": "B", "Ti": "C"},
    )
    loaded = load_config(destination)
    assert loaded.profile.distance_column == "A"
    assert loaded.profile.composition_columns == {"Ni": "B", "Ti": "C"}

    with pytest.raises(ConfigurationError, match="duplicated"):
        write_starter_config(
            workbook_path,
            tmp_path / "ambiguous.yaml",
            ["Ni=0.1", "Ti=0.2"],
            sheet="Gradient",
        )


def test_discovery_validation_errors(tmp_path: Path):
    profile = tmp_path / "profile.csv"
    profile.write_text("distance_nm,A\n0,1\n", encoding="utf-8")
    with pytest.raises(ProfileError, match="beyond the end"):
        inspect_profile_file(profile, header_row=99)
    unsupported = profile.with_suffix(".txt")
    unsupported.write_text("data", encoding="utf-8")
    with pytest.raises(ProfileError, match="Unsupported"):
        inspect_profile_file(unsupported)
    with pytest.raises(ConfigurationError, match="At least one"):
        write_starter_config(profile, tmp_path / "config.yaml", [])
    with pytest.raises(ConfigurationError, match="unique"):
        write_starter_config(profile, tmp_path / "config.yaml", ["A=1", "a=2"])
