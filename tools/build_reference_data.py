"""Build compact regression fixtures from an exported MATLAB baseline.

This utility reads but never modifies the original MATLAB script or workbook.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mat", required=True, type=Path)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    data = loadmat(args.mat)
    arrays = {
        key: np.asarray(data[key]).squeeze()
        for key in (
            "L",
            "x",
            "k_solved",
            "dt",
            "t",
            "power",
            "POWER",
            "xtest",
            "ImportantIndices",
            "K",
        )
    }
    np.savez_compressed(args.output / "matlab_baseline.npz", **arrays)

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["Distance Optimized"]
    with (args.output / "matlab_profile.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["distance_nm", "Ni", "Ti", "Nb"])
        for values in sheet.iter_rows(
            min_row=3, max_row=603, min_col=1, max_col=4, values_only=True
        ):
            writer.writerow(values)
    workbook.close()

    runtime = float(np.asarray(data["matlab_runtime_seconds"]).squeeze())
    metadata = {
        "matlab_version": "R2024b",
        "matlab_runtime_seconds": runtime,
        "source_file": args.source.name,
        "source_sha256": sha256(args.source),
        "workbook_file": args.workbook.name,
        "workbook_sha256": sha256(args.workbook),
        "profile_sheet": "Distance Optimized",
        "profile_range": "A3:D603",
        "target_rates_nm_per_min_per_watt": [0.1003, 0.0533, 0.0710],
        "total_power_w": 300.0,
    }
    (args.output / "baseline_metadata.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
